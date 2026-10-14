#!/usr/bin/env python3
"""Word文档美化脚本：目录、标题、正文样式"""
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# 颜色定义
DARK_BLUE = RGBColor(0x1F, 0x3A, 0x5F)
MEDIUM_BLUE = RGBColor(0x2E, 0x5A, 0x88)
LIGHT_BLUE = RGBColor(0x3A, 0x6B, 0x9E)
LIGHT_GRAY = RGBColor(0x55, 0x55, 0x55)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)

def add_bottom_border(paragraph, color='1F3A5F', size='8'):
    pPr = paragraph._element.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), size)
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)

def add_left_border(paragraph, color='2E5A88', size='18'):
    pPr = paragraph._element.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), size)
    left.set(qn('w:space'), '8')
    left.set(qn('w:color'), color)
    pBdr.append(left)
    pPr.append(pBdr)

def beautify_document(docx_path):
    doc = Document(docx_path)
    
    # 1. 美化标题
    for p in doc.paragraphs:
        if p.style.name == 'Heading 1':
            for run in p.runs:
                run.font.name = '黑体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
                run.font.size = Pt(18)
                run.font.bold = True
                run.font.color.rgb = DARK_BLUE
            p.paragraph_format.space_before = Pt(24)
            p.paragraph_format.space_after = Pt(12)
            add_bottom_border(p)
        
        elif p.style.name == 'Heading 2':
            for run in p.runs:
                run.font.name = '黑体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
                run.font.size = Pt(15)
                run.font.bold = True
                run.font.color.rgb = MEDIUM_BLUE
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
        
        elif p.style.name == 'Heading 3':
            for run in p.runs:
                run.font.name = '黑体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
                run.font.size = Pt(13)
                run.font.bold = True
                run.font.color.rgb = LIGHT_BLUE
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
    
    # 2. 美化正文
    for p in doc.paragraphs:
        if p.style.name == 'Normal':
            for run in p.runs:
                if run.font.color.rgb is None or run.font.color.rgb == RGBColor(0, 0, 0):
                    run.font.color.rgb = BLACK
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(6)
    
    # 3. 美化目录（找到目录区域）
    toc_start = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() in ('目  录', '目录'):
            toc_start = i
            break
    
    if toc_start is not None:
        title_p = doc.paragraphs[toc_start]
        for run in title_p.runs:
            run._element.getparent().remove(run._element)
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_p.paragraph_format.space_before = Pt(24)
        title_p.paragraph_format.space_after = Pt(20)
        run = title_p.add_run('目      录')
        run.font.size = Pt(26)
        run.font.bold = True
        run.font.color.rgb = DARK_BLUE
        run.font.name = '黑体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        add_bottom_border(title_p, size='12')
        
        # 目录条目
        for i in range(toc_start + 1, len(doc.paragraphs)):
            p = doc.paragraphs[i]
            if p.style.name == 'Heading 1':
                break
            text = p.text.strip()
            if not text:
                continue
            is_level1 = not text.startswith('  ')
            for run in p.runs:
                run._element.getparent().remove(run._element)
            if is_level1:
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.left_indent = Cm(0.5)
                run = p.add_run(text)
                run.font.size = Pt(14)
                run.font.bold = True
                run.font.color.rgb = MEDIUM_BLUE
                run.font.name = '黑体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
                add_left_border(p)
            else:
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.left_indent = Cm(2.0)
                run = p.add_run(text.strip())
                run.font.size = Pt(12)
                run.font.color.rgb = LIGHT_GRAY
                run.font.name = '宋体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    doc.save(docx_path)
    print(f"文档美化完成 → {docx_path}")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("用法: python beautify_docx.py output.docx")
        sys.exit(1)
    beautify_document(sys.argv[1])
