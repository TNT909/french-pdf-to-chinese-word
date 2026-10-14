#!/usr/bin/env python3
"""Word文档校验脚本：检查结构、表格、页码标签"""
import re
import sys
from docx import Document
from docx.oxml.ns import qn

def validate_document(docx_path):
    doc = Document(docx_path)
    issues = []
    
    # 1. 文档结构
    h1 = h2 = h3 = 0
    headings = []
    for p in doc.paragraphs:
        if p.style.name == 'Heading 1':
            h1 += 1
            headings.append(('H1', p.text))
        elif p.style.name == 'Heading 2':
            h2 += 1
            headings.append(('H2', p.text))
        elif p.style.name == 'Heading 3':
            h3 += 1
            headings.append(('H3', p.text))
    
    print(f"\n=== 文档结构 ===")
    print(f"H1: {h1}, H2: {h2}, H3: {h3}, 总计: {len(headings)}")
    
    # 2. 表格校验
    print(f"\n=== 表格校验 ===")
    for i, tbl in enumerate(doc.tables):
        rows = len(tbl.rows)
        cols = len(tbl.columns)
        print(f"表格{i+1}: {rows}行 x {cols}列")
        # 检查重复行
        for r_idx in range(1, rows):
            cells = [c.text.strip() for c in tbl.rows[r_idx].cells]
            prev_cells = [c.text.strip() for c in tbl.rows[r_idx-1].cells]
            if cells[0] and cells[0] and cells[0] in prev_cells[0] and cells[0] != '':
                # 检查是否为真重复（所有非空单元格相同）
                non_empty_curr = [c for c in cells if c]
                non_empty_prev = [c for c in prev_cells if c]
                if non_empty_curr and non_empty_prev and set(non_empty_curr) == set(non_empty_prev):
                    issues.append(f"表格{i+1}第{r_idx+1}行与第{r_idx}行重复: {cells[0][:40]}")
    
    # 3. 页码标签
    page_labels = []
    for p in doc.paragraphs:
        text = p.text.strip()
        m1 = re.match(r'\[原文页码 (\d+)\]', text)
        m2 = re.match(r'—— 原文第(\d+)页 ——', text)
        if m1:
            page_labels.append(int(m1.group(1)))
        elif m2:
            page_labels.append(int(m2.group(1)))
    
    print(f"\n=== 页码标签 ===")
    print(f"标签数: {len(page_labels)}")
    if page_labels:
        print(f"范围: {min(page_labels)}-{max(page_labels)}")
        print(f"覆盖页数: {len(set(page_labels))}")
    
    # 4. 法语残留检查（排除公司名和标准号）
    known_names = ['SONELGAZ', 'SONATRACH', 'DCQAV', 'Spa', 'STOS', 'ISO', 'EN', 'IEC',
                   'BNA', 'SWIFT', 'COPU', 'BAOSEM', 'PAQ', 'SCADA', 'CNAN',
                   'AGS', 'JdB', 'JDB', 'Sect', 'Plage', 'Embase', 'Cu', 'Al', 'mm']
    french_patterns = [
        r'\ble \b', r'\bla \b', r'\bles \b', r'\bdes \b', r'\bune \b',
        r'\best \b', r'\bsont \b', r'\bpour \b', r'\bdans \b', r'\bavec \b',
        r'\bpar \b', r'\bsur \b', r'\bqui \b', r'\bque \b', r'\baux \b'
    ]
    french_found = 0
    for p in doc.paragraphs:
        text = p.text
        # 跳过包含已知名称的段落
        if any(name in text for name in known_names):
            continue
        for pattern in french_patterns:
            if re.search(pattern, text.lower()):
                french_found += 1
                issues.append(f"法语残留: {text[:60]}...")
                break
    
    print(f"\n=== 法语残留 ===")
    print(f"发现 {french_found} 处")
    
    # 5. 标准编号保留检查
    standards = re.findall(r'(ISO\s?\d+|EN\s?\d+|IEC\s?\d+|N\.INT/DCQAV/\d+)', doc.element.xml)
    print(f"\n=== 标准编号 ===")
    print(f"保留的标准编号: {set(standards)}")
    
    # 汇总
    print(f"\n=== 校验结果 ===")
    if issues:
        print(f"发现 {len(issues)} 个问题:")
        for issue in issues[:20]:
            print(f"  ❌ {issue}")
        if len(issues) > 20:
            print(f"  ... 还有 {len(issues)-20} 个问题")
    else:
        print("✅ 校验通过，无问题")
    
    return issues

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("用法: python validate_docx.py output.docx")
        sys.exit(1)
    validate_document(sys.argv[1])
